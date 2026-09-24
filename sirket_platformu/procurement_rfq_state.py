import io
import re
import reflex as rx
import pandas as pd
import smtplib
import ssl
import imaplib
import time
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders

from typing import List, Dict, Any, Union
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

from .dashboard_state import DashboardState
from .shared_quotes import SHARED_QUOTES

try:
    from .services.db_service import get_all_teklifler_from_db
except ImportError:
    try:
        from services.db_service import get_all_teklifler_from_db
    except ImportError:
        get_all_teklifler_from_db = None


class ProcurementRfqState(rx.State):
    currency: str = "TRY (₺)"
    
    # Teklif Seçenekleri
    teklif_secenekleri: List[str] = []
    secilen_teklif: str = ""
    
    # Dahili sözlük haritası
    _quotes_map: Dict[str, Any] = {}
    
    # İhracat Formatı
    ihracat_formatlari: List[str] = [
        "Tedarikçi Teklif Talep Formu (Fiyatsız - RFQ)",
        "Birim Maliyetli Kontrol Listesi (İç Denetim)",
        "Toplu Satınalma İcmal Tablosu"
    ]
    secilen_format: str = "Tedarikçi Teklif Talep Formu (Fiyatsız - RFQ)"
    
    # E-Posta Simülatörü Alanları
    tedarikci_epostalar: str = "satis@vega.com, teklif@endress.com"
    son_yanit_tarihi: str = (datetime.now() + timedelta(days=1)).strftime("%d.%m.%Y")

    # Uzman Posta SMTP Yapılandırması
    smtp_server: str = "mail.uzmanposta.com"
    smtp_port: int = 465
    sender_email: str = "m.guner@petrotekelektrik.com.tr"
    sender_password: str = "Mu3@19Gu5"
    
    def set_currency(self, val: Any):
        if isinstance(val, list) and len(val) > 0:
            self.currency = str(val[0])
        else:
            self.currency = str(val)

    def set_secilen_format(self, val: str):
        self.secilen_format = val

    def set_tedarikci_epostalar(self, val: str):
        self.tedarikci_epostalar = val

    def set_son_yanit_tarihi(self, val: str):
        self.son_yanit_tarihi = val

    async def on_load(self):
        """DashboardState ve DB'deki teklifleri açılır listeye yükler."""
        await self.load_quotes()

    async def load_quotes(self):
        dash_state = await self.get_state(DashboardState)
        raw_list = getattr(dash_state, "raw_quotes", [])

        options = []
        self._quotes_map = {}

        for q in raw_list:
            kod = q.get("kod") or q.get("teklif_kodu", "")
            musteri = q.get("musteri", "-")
            if kod:
                label = f"{kod} | {musteri}"
                if label not in options:
                    options.append(label)
                    self._quotes_map[label] = q

        if get_all_teklifler_from_db:
            try:
                db_quotes = get_all_teklifler_from_db()
                for q in db_quotes:
                    kod = q.get("kod") or q.get("teklif_kodu", "")
                    musteri = q.get("musteri", "-")
                    if kod:
                        label = f"{kod} | {musteri}"
                        if label not in options:
                            options.append(label)
                            self._quotes_map[label] = q
            except Exception:
                pass

        if not options:
            options = ["PT202609201555 Deneme | Shell"]

        self.teklif_secenekleri = options
        if not self.secilen_teklif or self.secilen_teklif not in options:
            self.secilen_teklif = options[0]

    def set_secilen_teklif(self, val: str):
        self.secilen_teklif = val

    @rx.var
    def secilen_kod(self) -> str:
        if "|" in self.secilen_teklif:
            return self.secilen_teklif.split("|")[0].strip()
        return self.secilen_teklif.strip()

    @rx.var
    def ek_dosya_adi(self) -> str:
        clean_code = re.sub(r"[^\w\-]", "_", self.secilen_kod)
        return f"RFQ_{clean_code}_Tedarikci.xlsx"

    # Excel Üretim Metodu (Hem indirme hem mail için ortak)
    def generate_rfq_excel_bytes(self) -> io.BytesIO:
        """Tedarikçiye gidecek sarı dolgulu RFQ Excel dosyasını bellekte oluşturur."""
        kalemler = []
        if hasattr(self, "_quotes_map") and self.secilen_teklif in self._quotes_map:
            q = self._quotes_map[self.secilen_teklif]
            kalemler = q.get("kalemler", [])
        if not kalemler and self.secilen_kod in SHARED_QUOTES:
            kalemler = SHARED_QUOTES[self.secilen_kod].get("kalemler", [])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tedarikçi Teklif Talep Formu"

        header_fill = PatternFill(start_color="0A1020", end_color="0A1020", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        yellow_fill = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.merge_cells("A1:G1")
        ws["A1"] = f"MALZEME / HİZMET TEKLİF TALEP FORMU (RFQ) - {self.secilen_kod}"
        ws["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="0F172A")
        ws["A1"].alignment = Alignment(vertical="center")

        ws["A2"] = f"Son Teklif Verme Tarihi: {self.son_yanit_tarihi}"
        ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="475569")

        headers = [
            "Sıra No",
            "Malzeme / Hizmet Tanımı",
            "Miktar",
            "Birim",
            "Tedarikçi Birim Fiyatı (Doldurunuz)",
            "Para Birimi",
            "Teslim Süresi (Hafta / Gün)"
        ]

        row_idx = 4
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 3, 4, 6, 7] else "left", vertical="center")
            cell.border = thin_border

        if kalemler:
            for i, it in enumerate(kalemler, start=1):
                r = row_idx + i
                ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
                ws.cell(row=r, column=2, value=it.get("malzeme_adi") or it.get("tanim", "Kalem")).alignment = Alignment(horizontal="left")
                ws.cell(row=r, column=3, value=float(it.get("miktar", 1.0))).alignment = Alignment(horizontal="center")
                ws.cell(row=r, column=4, value=str(it.get("birim", "Adet"))).alignment = Alignment(horizontal="center")
                
                c_price = ws.cell(row=r, column=5)
                c_price.fill = yellow_fill
                c_price.border = thin_border

                c_curr = ws.cell(row=r, column=6, value="EUR")
                c_curr.fill = yellow_fill
                c_curr.alignment = Alignment(horizontal="center")
                c_curr.border = thin_border

                c_lead = ws.cell(row=r, column=7)
                c_lead.fill = yellow_fill
                c_lead.border = thin_border

                for c in range(1, 5):
                    ws.cell(row=r, column=c).border = thin_border
        else:
            r = row_idx + 1
            ws.cell(row=r, column=1, value=1).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=2, value=f"{self.secilen_kod} Kapsamı Malzeme Paketi")
            ws.cell(row=r, column=3, value=1).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=4, value="Set").alignment = Alignment(horizontal="center")
            for c in range(5, 8):
                cell = ws.cell(row=r, column=c)
                cell.fill = yellow_fill
                cell.border = thin_border

        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 45
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 30
        ws.column_dimensions['F'].width = 14
        ws.column_dimensions['G'].width = 25

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    # Eski adı çağıran yerler için geriye dönük uyumluluk takma adı (alias)
    def _generate_rfq_excel_bytes(self) -> io.BytesIO:
        return self.generate_rfq_excel_bytes()

    async def export_rfq_excel(self):
        """Tedarikçiye gönderilecek standart sarı dolgulu RFQ Excel tablosunu üretir ve indirir."""
        buffer = self.generate_rfq_excel_bytes()
        return rx.download(
            data=buffer.getvalue(),
            filename=self.ek_dosya_adi
        )

    async def send_supplier_email(self):
        """Uzman Posta SSL (Port 465) üzerinden RFQ gönderir ve Outlook Gönderilenler klasörüne kaydeder."""
        if not self.sender_password.strip():
            return rx.toast.error("Lütfen posta kutusu şifrenizi girin.", position="top-right")

        if not self.tedarikci_epostalar.strip():
            return rx.toast.error("Lütfen en az bir tedarikçi e-posta adresi girin.", position="top-right")

        recipients = [
            e.strip()
            for e in self.tedarikci_epostalar.replace(";", ",").split(",")
            if e.strip()
        ]

        if not recipients:
            return rx.toast.error("Geçerli bir tedarikçi e-posta adresi bulunamadı.", position="top-right")

        excel_buffer = self.generate_rfq_excel_bytes()
        excel_bytes = excel_buffer.getvalue()

        # Modern Kurumsal HTML E-Posta Şablonu
        # Outlook (Word motoru) ve Gmail tam uyumlu Tablo Tabanlı HTML Şablonu
        body_html = f"""<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.0 Transitional//EN" "http://www.w3.org/TR/xhtml1/DTD/xhtml1-transitional.dtd">
            <html xmlns="http://www.w3.org/1999/xhtml">
            <head>
                <meta http-equiv="Content-Type" content="text/html; charset=UTF-8" />
                <meta name="viewport" content="width=device-width, initial-scale=1.0"/>
                <title>PetroTek RFQ</title>
            </head>
            <body style="margin: 0; padding: 20px 0; background-color: #f1f5f9; font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; -webkit-text-size-adjust: 100%; -ms-text-size-adjust: 100%;">
                <!--[if mso]>
                <table align="center" border="0" cellspacing="0" cellpadding="0" width="600">
                <tr>
                <td>
                <![endif]-->
                <table align="center" border="0" cellpadding="0" cellspacing="0" width="100%" style="max-width: 600px; margin: 0 auto; background-color: #ffffff; border: 1px solid #e2e8f0;">
                    <!-- Üst Başlık (Lacivert Zemin) -->
                    <tr>
                        <td bgcolor="#0f172a" style="padding: 26px 30px; background-color: #0f172a; border-bottom: 3px solid #0284c7;">
                            <table border="0" cellpadding="0" cellspacing="0" width="100%">
                                <tr>
                                    <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 20px; font-weight: bold; color: #ffffff; line-height: 24px;">
                                        PetroTek Engineering
                                    </td>
                                </tr>
                                <tr>
                                    <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 13px; color: #94a3b8; padding-top: 4px; line-height: 18px;">
                                        Satınalma &amp; Tedarik Yönetim Departmanı
                                    </td>
                                </tr>
                            </table>
                        </td>
                    </tr>
                    
                    <!-- İçerik Alanı -->
                    <tr>
                        <td style="padding: 28px 30px;">
                            <table border="0" cellpadding="0" cellspacing="0" width="100%">
                                <tr>
                                    <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 14px; font-weight: bold; color: #0f172a; padding-bottom: 12px;">
                                        Sayın Yetkili,
                                    </td>
                                </tr>
                                <tr>
                                    <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 13.5px; line-height: 20px; color: #334155; padding-bottom: 20px;">
                                        Aşağıda proje kodu belirtilen iş kapsamımızda ihtiyaç duyulan malzeme, ekipman ve hizmet kalemleri için birim fiyat ve teslim süresi teklifinizi rica ederiz.
                                    </td>
                                </tr>
                                
                                <!-- Proje Bilgi Kartı -->
                                <tr>
                                    <td style="padding-bottom: 20px;">
                                        <table border="0" cellpadding="10" cellspacing="0" width="100%" style="background-color: #f8fafc; border: 1px solid #e2e8f0;">
                                            <tr>
                                                <td width="40%" style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 12.5px; color: #64748b; border-bottom: 1px solid #edf2f7;">
                                                    Proje / Teklif Kodu:
                                                </td>
                                                <td width="60%" style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 13px; font-weight: bold; color: #0f172a; border-bottom: 1px solid #edf2f7;">
                                                    {self.secilen_kod}
                                                </td>
                                            </tr>
                                            <tr>
                                                <td width="40%" style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 12.5px; color: #64748b; border-bottom: 1px solid #edf2f7;">
                                                    Son Teklif Verme Tarihi:
                                                </td>
                                                <td width="60%" style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 13px; font-weight: bold; color: #dc2626; border-bottom: 1px solid #edf2f7;">
                                                    {self.son_yanit_tarihi}
                                                </td>
                                            </tr>
                                            <tr>
                                                <td width="40%" style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 12.5px; color: #64748b;">
                                                    Talep Türü:
                                                </td>
                                                <td width="60%" style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 12.5px; color: #0f172a; font-weight: 600;">
                                                    Tedarikçi Teklif Talep Formu (RFQ)
                                                </td>
                                            </tr>
                                        </table>
                                    </td>
                                </tr>
                                
                                <!-- Sarı Uyarı Kutusu -->
                                <tr>
                                    <td style="padding-bottom: 20px;">
                                        <table border="0" cellpadding="12" cellspacing="0" width="100%" style="background-color: #fefce8; border-left: 4px solid #eab308; border-top: 1px solid #fef08a; border-right: 1px solid #fef08a; border-bottom: 1px solid #fef08a;">
                                            <tr>
                                                <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 12.5px; line-height: 18px; color: #854d0e;">
                                                    <strong>Önemli Not:</strong> Lütfen ekte yer alan Excel dosyasındaki <strong>sarı dolgulu</strong> hücreleri (birim fiyat, para birimi ve temin süresi) eksiksiz doldurarak bu e-postayı yanıtlayınız.
                                                </td>
                                            </tr>
                                        </table>
                                    </td>
                                </tr>
                                
                                <!-- Ek Dosya Rozeti -->
                                <tr>
                                    <td style="padding-bottom: 8px;">
                                        <table border="0" cellpadding="6" cellspacing="0" style="background-color: #e0f2fe; border: 1px solid #bae6fd;">
                                            <tr>
                                                <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 12px; font-weight: bold; color: #0369a1;">
                                                    &#128206; Ek: {self.ek_dosya_adi}
                                                </td>
                                            </tr>
                                        </table>
                                    </td>
                                </tr>
                            </table>
                        </td>
                    </tr>
                    
                    <!-- Alt İmza Alanı -->
                    <tr>
                        <td bgcolor="#f8fafc" style="padding: 18px 30px; background-color: #f8fafc; border-top: 1px solid #e2e8f0;">
                            <table border="0" cellpadding="0" cellspacing="0" width="100%">
                                <tr>
                                    <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 12px; font-weight: bold; color: #0f172a; line-height: 16px;">
                                        PetroTek Elektrik &amp; Mühendislik Ltd. Şti.
                                    </td>
                                </tr>
                                <tr>
                                    <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 11.5px; color: #64748b; line-height: 16px; padding-top: 2px;">
                                        Endüstriyel Otomasyon &amp; Petrol Tesisleri Çözümleri
                                    </td>
                                </tr>
                                <tr>
                                    <td style="font-family: 'Segoe UI', Arial, sans-serif; font-size: 11.5px; color: #64748b; line-height: 16px; padding-top: 2px;">
                                        E-Posta: <a href="mailto:{self.sender_email}" style="color: #0284c7; text-decoration: none;">{self.sender_email}</a>
                                    </td>
                                </tr>
                            </table>
                        </td>
                    </tr>
                </table>
                <!--[if mso]>
                </td>
                </tr>
                </table>
                <![endif]-->
            </body>
            </html>"""

        subject = f"Teklif Talebi (RFQ) - {self.secilen_kod} - Son Tarih: {self.son_yanit_tarihi}"

        try:
            msg = MIMEMultipart()
            msg["From"] = self.sender_email
            msg["To"] = ", ".join(recipients)
            msg["Subject"] = subject
            msg["Bcc"] = self.sender_email.strip()
            msg.attach(MIMEText(body_html, "html", "utf-8"))

            part = MIMEBase("application", "octet-stream")
            part.set_payload(excel_bytes)
            encoders.encode_base64(part)
            part.add_header("Content-Disposition", f'attachment; filename="{self.ek_dosya_adi}"')
            msg.attach(part)

            context = ssl.create_default_context()
            context.check_hostname = False
            context.verify_mode = ssl.CERT_NONE

            all_send_recipients = list(set(recipients + [self.sender_email.strip()]))

            # 1. SMTP ile İletim (SSL 465)
            with smtplib.SMTP_SSL("mail.uzmanposta.com", 465, timeout=25, context=context) as server:
                server.login(self.sender_email.strip(), self.sender_password.strip())
                server.sendmail(self.sender_email.strip(), all_send_recipients, msg.as_string())

            # 2. IMAP Üzerinden Sent/Gönderilenler Kutusuna Kaydetme (SSL 993)
            imap_bilgi = ""
            try:
                with imaplib.IMAP4_SSL("mail.uzmanposta.com", 993, ssl_context=context) as imap:
                    imap.login(self.sender_email.strip(), self.sender_password.strip())
                    
                    typ, folder_list = imap.list()
                    secilen_klasor_raw = None

                    if folder_list:
                        for f in folder_list:
                            f_str = f.decode("latin1", errors="ignore")
                            if "\\Sent" in f_str:
                                parts = f_str.split(' "." ') if ' "." ' in f_str else f_str.split(' "/" ')
                                secilen_klasor_raw = parts[-1].strip().strip('"')
                                break
                        
                        if not secilen_klasor_raw:
                            for f in folder_list:
                                f_str = f.decode("latin1", errors="ignore")
                                for keyword in ["sent", "gonderil", "gönderil", "&bb8-nderilmi&ba8-"]:
                                    if keyword in f_str.lower():
                                        parts = f_str.split(' "." ') if ' "." ' in f_str else f_str.split(' "/" ')
                                        secilen_klasor_raw = parts[-1].strip().strip('"')
                                        break
                                if secilen_klasor_raw:
                                    break

                    aday_klasorler = []
                    if secilen_klasor_raw:
                        aday_klasorler.append(secilen_klasor_raw)
                    aday_klasorler.extend([
                        "INBOX.Sent",
                        "Sent",
                        "INBOX.&BB8-nderilmi&BA8- &ANY-geler",
                        "INBOX.Gonderilenler",
                        "Sent Items"
                    ])

                    kaydedildi = False
                    basarili_klasor = ""
                    for k in aday_klasorler:
                        try:
                            k_param = f'"{k}"' if not k.startswith('"') else k
                            res, resp_data = imap.append(
                                k_param,
                                "\\Seen",
                                imaplib.Time2Internaldate(time.time()),
                                msg.as_bytes()
                            )
                            if res == "OK":
                                kaydedildi = True
                                basarili_klasor = k
                                imap.select(k_param)
                                imap.noop()
                                break
                        except Exception:
                            continue

                    if kaydedildi:
                        imap_bilgi = f" (Kayıt Yeri: {basarili_klasor})"
                    else:
                        imap_bilgi = " (Sunucu Sent klasörü eşleşmedi)"

            except Exception as e:
                imap_bilgi = f" (IMAP Hatası: {str(e)})"

            return rx.toast.success(
                f"RFQ iletildi ve kayda alındı!{imap_bilgi}",
                position="top-right"
            )

        except smtplib.SMTPAuthenticationError:
            return rx.toast.error("SMTP Giriş Hatası: E-posta adresi veya şifre yanlış.", position="top-right")
        except Exception as e:
            return rx.toast.error(f"E-posta gönderilirken hata oluştu: {str(e)}", position="top-right")