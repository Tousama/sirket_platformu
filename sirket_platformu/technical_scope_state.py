import io
import os
import json
import re
import reflex as rx
from typing import List, Dict, Any, Tuple
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from .services.db_service import get_all_teklifler_from_db

MEMORY_FILE = "nlp_feedback_memory.json"


class TechnicalScopeState(rx.State):
    referans_secenekleri: List[str] = []
    secilen_referans: str = ""
    sartname_no: str = ""
    muhatap_hitap: str = ""

    giris_yazisi: str = ""
    is_kapsami: str = ""
    isveren_sorumluluklari: str = ""
    haric_tutulanlar: str = ""

    # Geri Bildirim ve Durum Değişkenleri
    is_analyzing: bool = False
    uploaded_file_name: str = ""
    tespit_edilen_tip: str = ""
    feedback_notes: str = ""
    saved_feedback_count: int = 0

    # Setters
    def set_secilen_referans(self, val: str):
        self.secilen_referans = val

    def set_sartname_no(self, val: str):
        self.sartname_no = val

    def set_muhatap_hitap(self, val: str):
        self.muhatap_hitap = val

    def set_giris_yazisi(self, val: str):
        self.giris_yazisi = val

    def set_is_kapsami(self, val: str):
        self.is_kapsami = val

    def set_isveren_sorumluluklari(self, val: str):
        self.isveren_sorumluluklari = val

    def set_haric_tutulanlar(self, val: str):
        self.haric_tutulanlar = val

    def set_feedback_notes(self, val: str):
        self.feedback_notes = val

    # =========================================================================
    # VERİTABANI VE İLK YÜKLEME
    # =========================================================================
    async def on_load(self):
        """Sayfa açıldığında veritabanındaki teklifleri ve hafızayı yükler."""
        await self.load_reference_quotes()
        memory = self._load_memory()
        self.saved_feedback_count = len(memory)
        
    async def load_reference_quotes(self):
        try:
            db_quotes = get_all_teklifler_from_db()
            options = []
            for q in db_quotes:
                kod = q.get("kod", "")
                musteri = q.get("musteri", "-")
                konu = q.get("konu", "-")
                options.append(f"{kod} | {musteri} | {konu}")
            self.referans_secenekleri = options if options else ["Genel Teklif Referansı"]
            if self.referans_secenekleri and not self.secilen_referans:
                self.secilen_referans = self.referans_secenekleri[0]
        except Exception:
            self.referans_secenekleri = ["Genel Teklif Referansı"]

    # =========================================================================
    # GERİ BİLDİRİM & ÖĞRENME MOTORU
    # =========================================================================
    def _load_memory(self) -> List[Dict[str, Any]]:
        if os.path.exists(MEMORY_FILE):
            try:
                with open(MEMORY_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                return []
        return []

    def _save_memory_entry(self, entry: Dict[str, Any]):
        memory = self._load_memory()
        memory.append(entry)
        with open(MEMORY_FILE, "w", encoding="utf-8") as f:
            json.dump(memory, f, ensure_ascii=False, indent=2)

    def save_user_feedback(self):
        entry = {
            "sartname_no": self.sartname_no,
            "uploaded_file": self.uploaded_file_name,
            "tespit_tipi": self.tespit_edilen_tip,
            "muhatap_hitap": self.muhatap_hitap,
            "approved_giris": self.giris_yazisi,
            "approved_scope": self.is_kapsami,
            "approved_employer": self.isveren_sorumluluklari,
            "approved_exclusions": self.haric_tutulanlar,
            "feedback_notes": self.feedback_notes,
        }
        self._save_memory_entry(entry)
        self.feedback_notes = ""
        memory = self._load_memory()
        self.saved_feedback_count = len(memory)
        return rx.toast.success(
            "Tercihler başarıyla NLP hafızasına kaydedildi!",
            position="top-right"
        )

    # =========================================================================
    # ŞARTNAME BİLGİ AYRIŞTIRMA (METADATA EXTRACTION)
    # =========================================================================
    def _extract_metadata(self, text: str) -> Tuple[str, str]:
        """Şartname no, revizyon ve müşteri unvanını tespit eder."""
        first_page = text[:2500]

        # Şartname No Tespiti (Örn: TS-GA-DOL-004-R00)
        ts_match = re.search(r"(?:TS\s*No|Teknik\s*Şartname\s*No)\s*[:|]?\s*([A-Z0-9\-_]+)", first_page, re.IGNORECASE)
        sartname_kod = ts_match.group(1).strip() if ts_match else ""

        # Başlık ve Konu Tespiti
        konu_match = re.search(r"KONU\s*\|\s*([^\n\r]+)", first_page, re.IGNORECASE)
        konu_str = konu_match.group(1).strip() if konu_match else ""

        # Müşteri ve Tesis Tespiti
        musteri_unvani = "Sayın Yetkili"
        if re.search(r"g[üu]zel\s*enerj[iı]", first_page, re.IGNORECASE):
            musteri_unvani = "Güzel Enerji Akaryakıt A.Ş. - Teknik Müdürlük / İkmal ve Lojistik Direktörlüğü"
            if not konu_str and "gebze" in first_page.lower():
                konu_str = "Gebze Akaryakıt Terminali Dolum Adası Revizyonu"
        elif re.search(r"t[üu]pra[sş]", first_page, re.IGNORECASE):
            musteri_unvani = "TÜPRAŞ Türkiye Petrol Rafinerileri A.Ş. - Bakım & Proje Müdürlüğü"
        elif re.search(r"shell", first_page, re.IGNORECASE):
            musteri_unvani = "Shell & Turcas Petrol A.Ş. - Mühendislik ve Operasyonlar Direktörlüğü"
        elif re.search(r"petrol\s*of[iı]s[iı]", first_page, re.IGNORECASE):
            musteri_unvani = "Petrol Ofisi A.Ş. - Terminal Mühendislik Departmanı"
        elif re.search(r"aves", first_page, re.IGNORECASE):
            musteri_unvani = "AVES Enerji Yağ ve Gıda Sanayi A.Ş. - Teknik İşler Direktörlüğü"

        # Şartname No Birleştirme
        if sartname_kod and konu_str:
            sartname_tam = f"{sartname_kod} - {konu_str}"
        elif sartname_kod:
            sartname_tam = sartname_kod
        elif konu_str:
            sartname_tam = konu_str
        else:
            sartname_tam = "Endüstriyel Tesis Teknik Şartnamesi"

        return sartname_tam, musteri_unvani

    # =========================================================================
    # ENDÜSTRİYEL KAPSAM VE SAHA NLP ÇÖZÜMLEME MOTORU
    # =========================================================================
    def _nlp_parse_specification(self, text: str) -> Dict[str, Any]:
        t_low = text.lower()

        # 1. Hafıza Kontrolü (Memory Cache Lookup)
        memory = self._load_memory()
        for item in reversed(memory):
            if item.get("sartname_no") and len(item["sartname_no"]) > 5:
                ref_key = item["sartname_no"].split("-")[0].strip().lower()
                if ref_key in t_low:
                    return {
                        "kapsam_tipi": item.get("tespit_tipi", "Hafızadan Alınan Şablon"),
                        "sartname_no": item.get("sartname_no", ""),
                        "muhatap": item.get("muhatap_hitap", "Sayın Yetkili"),
                        "giris": item.get("approved_giris", self.giris_yazisi),
                        "is_kapsami": item.get("approved_scope", ""),
                        "isveren": item.get("approved_employer", ""),
                        "haric": item.get("approved_exclusions", ""),
                    }

        sartname_no_bulunan, muhatap_bulunan = self._extract_metadata(text)

        # 2. Modül & Tesis Tipolojisi Skorlaması
        is_terminal_ada_revizyon = any(x in t_low for x in ["dolum adası", "ada 5", "ada-5", "peron", "junction box", "flashtech", "accuload", "katık"])
        is_tank_ciftligi = any(x in t_low for x in ["tank radar", "rex", "overfill", "scully", "tank sahası", "seviye şalteri"])

        montaj_keywords = [r"\bdemontaj\b", r"\bmontaj\b", r"\bkablaj\b", r"\bkablo\s*çekim", r"\btava\b", r"\brakor\b", r"\bsonlandırma\b", r"\bdevreye\s*alma\b", r"\bişçilik\b"]
        temin_keywords = [r"\bmalzeme\s*temin", r"\benstrüman\s*temin", r"\bsatın\s*alma\b", r"\btedarik\b", r"\bsupply\s*only\b"]

        montaj_skoru = sum(len(re.findall(k, t_low)) for k in montaj_keywords)
        temin_skoru = sum(len(re.findall(k, t_low)) for k in temin_keywords)

        is_kapsami_list = []
        isveren_list = []
        haric_list = []

        # SENARYO 1: AKARYAKIT TERMİNALİ DOLUM ADASI & KABLAJ REVİZYONU
        if is_terminal_ada_revizyon:
            kapsam_tipi = "Akaryakıt Terminali Dolum Adası & Elektrik-Kablaj Revizyonu (Anahtar Teslim)"
            ada_no = "5 No'lu (9-10 numaralı peron)" if ("ada 5" in t_low or "ada-5" in t_low) else "Dolum Adası"
            
            giris = (
                f"İşbu teknik teklif ve kapsam dokümanı; {muhatap_bulunan} bünyesindeki terminalde bulunan "
                f"{ada_no} alttan tanker dolum adasının elektrik bileşenleri, Junction Box'ları ve kablo tesisatının "
                f"yenilenmesi, saha testleri ve Terminal Otomasyon Sistemi (TAS) entegrasyonu ile anahtar teslim devreye alınmasını kapsamaktadır."
            )

            is_kapsami_list.append(f"• {ada_no} içerisindeki mevcut deforme kabloların, kablo tavalarının ve eski AC/DC bağlantı kutularının emniyetli demontajı ve terminal deposuna teslimi.")
            
            jb_match = re.search(r"(\d+)\s*adet\s*(?:yeni)?\s*junction\s*box", t_low)
            jb_adet = jb_match.group(1) if jb_match else "2"
            is_kapsami_list.append(f"• Saha koşullarına ve Ex-proof normlara tam uyumlu {jb_adet} adet yeni Junction Box temini, kaide montajı ve etiketlenmesi.")

            is_kapsami_list.append("• Şartnameye uygun zırhlı enerji kabloları, enstrüman sinyal kabloları, AccuLoad multi kabloları ve Ex-d/Ex-e zırhlı kablo glandlerinin (rakorlarının) temini ve çekimi.")
            is_kapsami_list.append("• Sıcak daldırma galvaniz kablo tavaları, kapakları, ek parçaları ve ağır hizmet mekanik support (destek) elemanlarının temini ve montajı.")

            if "katık" in t_low or "enjektör" in t_low:
                is_kapsami_list.append("• Mevcut katık enjektörleri ve ilgili saha enstrümanlarının elektriksel bağlantılarının kontrolü, yeni kablo tesisatına sonlandırılması ve test edilmesi.")

            if "flashtech" in t_low:
                is_kapsami_list.append("• Flashtech Terminal Otomasyon Sistemi ile entegrasyonun sağlanması, I/O adresleme ve dolum adası tanımlamalarının yüklenici mühendislerince yapılması.")
            else:
                is_kapsami_list.append("• Mevcut Terminal Otomasyon Sistemi (TAS) ile tam entegrasyon, sinyal testleri ve haberleşme doğrulamasının yapılması.")

            is_kapsami_list.append("• Saha kablo meger (izolasyon) testleri, pulse/sinyal süreklilik testleri, mühürleme ve sistemin eksiksiz devreye alınması.")
            is_kapsami_list.append("• İş bitiminde terminal personeline operasyonel eğitim verilmesi; As-Built projeler, test formları ve Kalite Kontrol Dosyasının (Soft & Hard Copy) teslimi.")

            isveren_list.append("• Çalışma yapılacak dolum adası ve hatların elektriksel izolasyonunun (LOTO) İşveren tarafından eksiksiz sağlanması.")
            if "forklift" in t_low:
                isveren_list.append("• Terminal sınırları içerisindeki ağır taşıma ve kaldırma işlerinde terminalin mevcut forkliftinin yüklenici kullanımına tahsisi.")
            isveren_list.append("• Demonte edilen malzemelerin istifleneceği terminal içi uygun depo alanının gösterilmesi.")
            isveren_list.append("• İSG sıcak/soğuk saha çalışma izinlerinin (Permit to Work) iş takvimini aksatmayacak şekilde onaylanması.")

            haric_list.append("• İnşaat, betonarme kaide, saha asfalt/zemin kırım ve hafriyat işleri.")
            haric_list.append("• Şartname kapsamında yer almayan mekanik borulama, boru deplasman ve kaynaklı hat tadilatları.")
            haric_list.append("• Ana otomasyon sunucuları ve Flashtech lisans ücretleri.")
            haric_list.append("• Terminal forklifti haricinde doğabilecek özel tonajlı vinç ve sepetli platform ihtiyaçları (gerektiğinde İşveren koordinasyonuyla sağlanır).")

        # SENARYO 2: TANK ÇİFTLİĞİ ENSTRÜMANTASYON & SEVİYE SİSTEMLERİ
        elif is_tank_ciftligi:
            kapsam_tipi = "Tank Sahası Enstrümantasyon & Aşırı Dolum Önleme (Overfill) Sistemi"
            giris = (
                f"İşbu teknik teklif dokümanı; {muhatap_bulunan} tesislerindeki depolama tanklarının "
                f"enstrümantasyon sistemlerinin montajı, tank üstü radar/şalter bağlantıları ve otomasyon odası entegrasyonunu kapsamaktadır."
            )
            is_kapsami_list.append("• Tank üstü seviye radarları, titreşimli çatal seviye şalterleri ve sıcaklık sensörlerinin montajı.")
            is_kapsami_list.append("• Ex-proof Scully / Topraklama ve Aşırı Dolum Önleme sistemlerinin kablajı ve interlock testleri.")
            is_kapsami_list.append("• RS-485 / Modbus haberleşme altyapısının çekimi ve PLC kontrol panelleri ile haberleştirilmesi.")
            is_kapsami_list.append("• Fabrika kalibrasyon sertifikaları, loop kontrol raporları ve devreye alma tutanaklarının teslimi.")

            isveren_list.append("• Tankların gazdan arındırılmış (gas-free) halde çalışmaya hazır teslim edilmesi.")
            isveren_list.append("• Tank flanş girişlerinin mekanik olarak montaja uygun hazır bulundurulması.")
            isveren_list.append("• İSG sıcak çalışma ve kapalı alan giriş izinlerinin zamanında sağlanması.")

            haric_list.append("• Tank nozul kaynakları, mekanik boru işleri ve yapısal çelik platform tadilatları.")
            haric_list.append("• Enerji temini için ana trafo/şalt sahasındaki ana pano tadilatları.")

        # SENARYO 3: SAF MALZEME / ENSTRÜMAN TEMİNİ (SUPPLY ONLY)
        elif montaj_skoru <= 2 or "sadece temin" in t_low:
            kapsam_tipi = "Endüstriyel Enstrüman & Malzeme Temini (Supply Only)"
            giris = (
                f"İşbu teknik teklif dokümanı; {muhatap_bulunan} tarafından talep edilen "
                f"şartname föylerine uygun enstrümanların orijinal fabrika çıkışlı olarak temini, sertifikasyonu ve şantiye teslimini kapsamaktadır."
            )
            is_kapsami_list.append("• Şartname teknik kriterlerine ve talep edilen ATEX/IECEx koruma sınıflarına haiz cihaz temini.")
            is_kapsami_list.append("• EN 10204 3.1 Malzeme İzlenebilirlik ve Fabrika Kalibrasyon Test Sertifikalarının teslimi.")
            is_kapsami_list.append("• Ekipmanların nakliye sigortalı (DAP/CIP) olarak tesis sahasına güvenli teslimatı.")

            isveren_list.append("• Malzemelerin şantiye sahasında teslim alınması, uygun kapalı depolama koşullarının sağlanması.")
            isveren_list.append("• Malzeme muayene ve kabul protokollerinin şartname takvimine uygun imzalanması.")

            haric_list.append("• Sahada mekanik ve elektriksel montaj, kablolama ve sonlandırma işçilikleri.")
            haric_list.append("• Saha loop testleri, enerji verme ve devreye alma hizmetleri.")

        # SENARYO 4: GENEL SAHA ELEKTRİK / ENSTRÜMANTASYON & MONTAJ
        else:
            kapsam_tipi = "Endüstriyel Elektrik, Kablaj & Enstrümantasyon Montajı"
            giris = (
                f"İşbu teknik teklif dokümanı; {muhatap_bulunan} bünyesinde gerçekleştirilecek "
                f"saha elektrik, kablo tavası montajı, kablo çekimi, sonlandırma ve enstrümantasyon devreye alma faaliyetlerini kapsamaktadır."
            )
            is_kapsami_list.append("• Saha kablo tavası güzergahlarının montajı, support imalatları ve zırhlı kablo çekimleri.")
            is_kapsami_list.append("• Ex-proof ekipman ve panolara ait kablo rakor bağlantıları, klemens sonlandırmaları ve etiketleme.")
            is_kapsami_list.append("• İzolasyon/meger testleri, 4-20mA HART sinyal doğrulama ve devreye alma testlerinin yürütülmesi.")
            is_kapsami_list.append("• Kırmızı hat (Red-line) / As-Built projelerin hazırlanması ve test dokümantasyonu teslimi.")

            isveren_list.append("• Sahada çalışma izinleri ve İSG sıcak çalışma formlarının zamanında onaylanması.")
            isveren_list.append("• Montaj öncesi proses hatlarının elektriksel ve operasyonel izolasyonunun sağlanması.")
            isveren_list.append("• Saha testleri ve çalışma süresince gerekli geçici şantiye enerjisi temini.")

            haric_list.append("• Her türlü inşaat, beton kaide ve saha hafriyat işleri.")
            haric_list.append("• Şartnamede açıkça belirtilmeyen mekanik borulama ve hat deplasmanları.")

        return {
            "kapsam_tipi": kapsam_tipi,
            "sartname_no": sartname_no_bulunan,
            "muhatap": muhatap_bulunan,
            "giris": giris,
            "is_kapsami": "\n".join(is_kapsami_list),
            "isveren": "\n".join(isveren_list),
            "haric": "\n".join(haric_list),
        }

    # =========================================================================
    # EVENT HANDLERS (PDF VE DOCX AYRIŞTIRICI)
    # =========================================================================
    async def handle_upload_and_analyze(self, files: List[rx.UploadFile]):
        if not files:
            yield rx.toast.error("Lütfen bir teknik şartname dosyası seçin.", position="top-right")
            return

        self.is_analyzing = True
        yield

        file = files[0]
        self.uploaded_file_name = file.filename
        content = await file.read()
        fn_low = file.filename.lower()

        raw_text = ""
        try:
            # 1. PDF İse:
            if fn_low.endswith(".pdf"):
                pdf_reader = PdfReader(io.BytesIO(content))
                for page in pdf_reader.pages:
                    t = page.extract_text()
                    if t:
                        raw_text += t + "\n"

            # 2. DOCX (Word) İse:
            elif fn_low.endswith(".docx"):
                doc = Document(io.BytesIO(content))
                # Paragrafları oku
                for p in doc.paragraphs:
                    if p.text.strip():
                        raw_text += p.text + "\n"
                # Tabloların içindeki metinleri ve şartname maddelerini oku
                for table in doc.tables:
                    for row in table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_vals:
                            raw_text += " | ".join(row_vals) + "\n"
            else:
                self.is_analyzing = False
                yield rx.toast.error("Yalnızca PDF ve DOCX dosyaları desteklenmektedir.", position="top-right")
                return

        except Exception as e:
            self.is_analyzing = False
            yield rx.toast.error(f"Dosya okunamadı: {str(e)}", position="top-right")
            return

        if not raw_text.strip():
            self.is_analyzing = False
            yield rx.toast.warning("Yüklenen dosya içerisinde okunabilir metin bulunamadı.", position="top-right")
            return

        parsed = self._nlp_parse_specification(raw_text)

        self.tespit_edilen_tip = parsed["kapsam_tipi"]
        self.sartname_no = parsed["sartname_no"]
        self.muhatap_hitap = parsed["muhatap"]
        self.giris_yazisi = parsed["giris"]
        self.is_kapsami = parsed["is_kapsami"]
        self.isveren_sorumluluklari = parsed["isveren"]
        self.haric_tutulanlar = parsed["haric"]

        self.is_analyzing = False
        yield rx.toast.success(
            f"Şartname başarıyla çözümlendi: '{self.tespit_edilen_tip}'",
            position="top-right"
        )

    def export_docx(self):
        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.9)
            section.right_margin = Inches(0.9)

        title_p = doc.add_paragraph()
        title_run = title_p.add_run("TEKNİK ŞARTNAME UYUMLULUK VE KAPSAM DOKÜMANI")
        title_run.font.name = "Calibri"
        title_run.font.size = Pt(15)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(14, 52, 98)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run("PetroTek Engineering - Endüstriyel Tesis & Otomasyon Çözümleri")
        sub_run.font.name = "Calibri"
        sub_run.font.size = Pt(9.5)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 116, 139)
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        table = doc.add_table(rows=3, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        info_pairs = [
            ("Referans Teklif Kodu / Proje:", self.secilen_referans),
            ("Teknik Şartname Referansı:", self.sartname_no),
            ("Muhatap / Müşteri:", self.muhatap_hitap),
        ]
        for idx, (label, val) in enumerate(info_pairs):
            row = table.rows[idx]
            lbl_cell = row.cells[0].paragraphs[0].add_run(label)
            lbl_cell.font.name = "Calibri"
            lbl_cell.font.size = Pt(9)
            lbl_cell.font.bold = True

            val_cell = row.cells[1].paragraphs[0].add_run(val)
            val_cell.font.name = "Calibri"
            val_cell.font.size = Pt(9)

        doc.add_paragraph()

        def add_section(heading_text: str, content_text: str):
            h_p = doc.add_paragraph()
            h_run = h_p.add_run(heading_text)
            h_run.font.name = "Calibri"
            h_run.font.size = Pt(11.5)
            h_run.font.bold = True
            h_run.font.color.rgb = RGBColor(30, 41, 59)

            c_p = doc.add_paragraph()
            c_run = c_p.add_run(content_text)
            c_run.font.name = "Calibri"
            c_run.font.size = Pt(9.5)
            c_p.paragraph_format.line_spacing = 1.2
            c_p.paragraph_format.space_after = Pt(8)

        add_section("1. GİRİŞ VE PROJE AMACI", self.giris_yazisi)
        add_section("2. İŞ KAPSAMI (SCOPE OF WORK)", self.is_kapsami)
        add_section("3. İŞVERENİN SORUMLULUKLARI", self.isveren_sorumluluklari)
        add_section("4. HARİÇ TUTULAN İŞLER (EXCLUSIONS)", self.haric_tutulanlar)

        doc.add_paragraph()
        sig_p = doc.add_paragraph()
        sig_run = sig_p.add_run("Teknik Onay & Hazırlayan:\nPetroTek Engineering Proje & Teklif Departmanı")
        sig_run.font.name = "Calibri"
        sig_run.font.size = Pt(9)
        sig_run.font.bold = True
        sig_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        target_stream = io.BytesIO()
        doc.save(target_stream)
        target_stream.seek(0)

        clean_ref = self.secilen_referans.split("|")[0].strip() if self.secilen_referans else "TEKLIF"
        return rx.download(
            data=target_stream.getvalue(),
            filename=f"Teknik_Kapsam_{clean_ref}.docx"
        )